from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_sample_assignment():
    doc = Document()

    # Header
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("DINAS PENDIDIKAN DAN KEBUDAYAAN\n")
    run.font.size = Pt(14)
    run.bold = True
    run = header.add_run("SMA NEGERI 1 EDULEARN\n")
    run.font.size = Pt(16)
    run.bold = True
    run = header.add_run("Jl. Teknologi No. 40, Jakarta Pusat\n")
    run.font.size = Pt(10)
    
    doc.add_paragraph("_" * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("\nTUGAS HARIAN: LOGIKA MATEMATIKA\n")
    run.font.size = Pt(12)
    run.bold = True

    # Student Info Section
    doc.add_paragraph("Nama\t: ............................")
    doc.add_paragraph("Kelas\t: ............................")
    doc.add_paragraph("No. Absen\t: ............................")
    doc.add_paragraph("\n")

    # Questions
    doc.add_paragraph("Kerjakan soal-soal di bawah ini dengan benar!", style='List Bullet')
    
    questions = [
        "Jelaskan apa yang dimaksud dengan pernyataan tunggal dan pernyataan majemuk dalam logika matematika!",
        "Tentukan tabel kebenaran untuk implikasi (P ∧ Q) → R.",
        "Berikan contoh pernyataan yang merupakan tautologi dalam kehidupan sehari-hari.",
        "Selesaikan penarikan kesimpulan menggunakan Modus Ponens dari premis: \n   - Jika hujan turun, maka jalanan basah. \n   - Hujan turun. \n   Kesimpulan: ...",
        "Buatlah satu contoh soal pilihan ganda mengenai negasi dari pernyataan berkuantor."
    ]

    for i, q in enumerate(questions, 1):
        p = doc.add_paragraph()
        p.add_run(f"{i}. {q}")

    doc.add_paragraph("\n\n--- Selamat Mengerjakan ---").alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Save the document
    file_path = "Contoh_Soal_Logika_Matematika.docx"
    doc.save(file_path)
    print(f"File created: {file_path}")

if __name__ == "__main__":
    create_sample_assignment()
