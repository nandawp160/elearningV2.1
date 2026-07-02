from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_bahasa_indonesia_assignment():
    doc = Document()

    # Header
    header = doc.add_paragraph()
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = header.add_run("DINAS PENDIDIKAN DAN KEBUDAYAAN\n")
    run.font.size = Pt(14)
    run.bold = True
    run = header.add_run("SMA NEGERI 1 EDULEARN\n")
    run.font.size = Pt(16)
    run.bold = True
    run = header.add_run("Jl. Teknologi No. 40, Jakarta Pusat\n")
    run.font.size = Pt(10)
    
    doc.add_paragraph("_" * 50).alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("\nULANGAN HARIAN BAB 1: TEKS LAPORAN HASIL OBSERVASI\n")
    run.font.size = Pt(12)
    run.bold = True

    # Student Info Section
    doc.add_paragraph("Nama\t: ............................")
    doc.add_paragraph("Kelas\t: ............................")
    doc.add_paragraph("\n")

    # Questions
    doc.add_paragraph("Jawablah pertanyaan di bawah ini dengan singkat dan jelas!", style='List Bullet')
    
    questions = [
        "Apa yang dimaksud dengan Teks Laporan Hasil Observasi (LHO) dan jelaskan tujuan utamanya!",
        "Sebutkan dan jelaskan struktur membangun Teks Laporan Hasil Observasi!"
    ]

    for i, q in enumerate(questions, 1):
        p = doc.add_paragraph()
        run = p.add_run(f"{i}. {q}")
        run.font.size = Pt(11)

    doc.add_paragraph("\n\n--- Selamat Mengerjakan ---").alignment = WD_ALIGN_PARAGRAPH.CENTER

    # Save the document
    file_path = "Latihan_Bahasa_Indonesia_Bab_1.docx"
    doc.save(file_path)
    print(f"File created: {file_path}")

if __name__ == "__main__":
    create_bahasa_indonesia_assignment()
