from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_exam_template():
    document = Document()

    # Style
    style = document.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(12)

    # Header
    header = document.sections[0].header
    p = header.paragraphs[0]
    p.text = "TEMPLATE SOAL UJIAN"
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    
    # Title
    heading = document.add_heading('JUDUL TUGAS / UJIAN', 0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    document.add_paragraph('Instruksi:')
    document.add_paragraph('1. Gunakan format ini untuk membuat soal.')
    document.add_paragraph('2. Untuk pilihan ganda, gunakan format A, B, C, D, E.')
    document.add_paragraph('3. Simpan file ini dan upload saat membuat tugas.')
    document.add_paragraph('')

    document.add_heading('BAGIAN 1: PILIHAN GANDA', level=1)
    
    # Example Question
    document.add_paragraph('1. Pertanyaan nomor 1 tuliskan disini...')
    document.add_paragraph('   A. Pilihan A')
    document.add_paragraph('   B. Pilihan B')
    document.add_paragraph('   C. Pilihan C')
    document.add_paragraph('   D. Pilihan D')
    document.add_paragraph('   E. Pilihan E')
    document.add_paragraph('')
    
    document.add_paragraph('2. Pertanyaan nomor 2 tuliskan disini...')
    document.add_paragraph('   A. Pilihan A')
    document.add_paragraph('   B. Pilihan B')
    document.add_paragraph('   C. Pilihan C')
    document.add_paragraph('   D. Pilihan D')
    document.add_paragraph('   E. Pilihan E')
    document.add_paragraph('')

    document.add_heading('BAGIAN 2: ESSAY', level=1)
    document.add_paragraph('1. Tuliskan pertanyaan essay nomor 1 disini...')
    document.add_paragraph('')
    document.add_paragraph('2. Tuliskan pertanyaan essay nomor 2 disini...')
    
    # Save
    import os
    if not os.path.exists('public/templates'):
        os.makedirs('public/templates')
        
    document.save('public/templates/template_soal.docx')
    print("Template created at public/templates/template_soal.docx")

if __name__ == "__main__":
    create_exam_template()
